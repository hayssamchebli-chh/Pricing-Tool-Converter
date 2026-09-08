"""Core logic: turn a "last purchases" export into the tendering Order workbook.

Reads a purchase-history file (one row per invoice line), keeps the three most
recent unit prices per item, and writes them into the Order template columns
``Last U.P.(1..3) (EUR)`` / ``Date(1..3)`` -- most recent first.
"""

from __future__ import annotations

import io
import re
from copy import copy
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.styles.numbers import is_date_format
from openpyxl.utils import get_column_letter

TEMPLATE_PATH = Path(__file__).with_name("template") / "order_template.xlsx"

MAX_HISTORY = 3

# Used when slot 1 of the template cannot supply a sane format to copy.
PRICE_FORMAT = "#,##0.00"
DATE_FORMAT = "mm-dd-yy"

# Header text is normalised (lower-case, letters and digits only) before
# matching, so "Last U.P.\n(1) (EUR)" and "InvLine_Unit_Cost" both resolve.
PURCHASE_ALIASES = {
    "item": ["no_", "no", "itemcode", "itemno", "itemnumber", "item", "code",
             "partnumber", "partno", "materialno", "material", "sku", "ref",
             "reference"],
    "description": ["description", "desc", "itemdescription", "itemname",
                    "name", "designation"],
    "date": ["invlinepostingdate", "postingdate", "invlinedate", "invoicedate",
             "documentdate", "purchasedate", "date"],
    "price": ["invlineunitcost", "unitcost", "unitprice", "unitpriceeur",
              "netprice", "lastup", "upeur", "price", "cost", "up"],
}

ORDER_TARGETS = {
    "item": ["itemcode", "itemno", "no_", "code"],
    "description": ["description"],
    "price1": ["lastup1eur", "lastup1"],
    "date1": ["date1"],
    "price2": ["lastup2eur", "lastup2"],
    "date2": ["date2"],
    "price3": ["lastup3eur", "lastup3"],
    "date3": ["date3"],
}


class BuildError(Exception):
    """Raised when an input file cannot be interpreted."""


def norm(value):
    """Normalise a header or key: 'Last U.P.\\n(1) (EUR)' -> 'lastup1eur'."""
    return re.sub(r"[^a-z0-9]", "", str(value).lower()) if value is not None else ""


def _resolve(columns, aliases):
    """Map logical names to actual column labels, preferring earlier aliases."""
    normalised = {}
    for col in columns:
        normalised.setdefault(norm(col), col)
    resolved = {}
    for logical, names in aliases.items():
        for name in names:
            if name in normalised:
                resolved[logical] = normalised[name]
                break
    return resolved


# --- reading the purchase history -----------------------------------------


def read_purchases(source, sheet_name=0):
    """Load a purchase-history file into a tidy frame: item, description, date, price."""
    name = getattr(source, "name", str(source)).lower()
    if name.endswith(".csv") or name.endswith(".txt"):
        frame = pd.read_csv(source, sep=None, engine="python")
    else:
        frame = pd.read_excel(source, sheet_name=sheet_name)

    columns = _resolve(frame.columns, PURCHASE_ALIASES)
    missing = [key for key in ("item", "date", "price") if key not in columns]
    if missing:
        raise BuildError(
            "Could not find the " + ", ".join(missing) + " column(s) in the "
            "uploaded file. Columns found: "
            + ", ".join(str(c) for c in frame.columns)
        )

    prices = frame[columns["price"]]
    if prices.dtype == object:
        prices = (
            prices.astype(str)
            .str.replace(r"[^\d.,\-]", "", regex=True)
            .str.replace(",", "", regex=False)
        )

    tidy = pd.DataFrame(
        {
            "item": frame[columns["item"]],
            "description": frame[columns["description"]] if "description" in columns else "",
            "date": pd.to_datetime(frame[columns["date"]], errors="coerce"),
            "price": pd.to_numeric(prices, errors="coerce"),
        }
    )
    tidy["item"] = tidy["item"].astype(str).str.strip()
    tidy["description"] = tidy["description"].fillna("").astype(str).str.strip()
    tidy["key"] = tidy["item"].str.upper()
    tidy["source_row"] = range(len(tidy))
    keep = (tidy["item"] != "") & (tidy["item"].str.lower() != "nan")
    return tidy.loc[keep].reset_index(drop=True)


@dataclass
class ItemHistory:
    """The three most recent purchases of one item, newest first."""

    item: str
    description: str = ""
    prices: list = field(default_factory=list)
    dates: list = field(default_factory=list)
    total_purchases: int = 0

    def cell(self, index):
        """Price/date pair at position `index` (0 = most recent), or (None, None)."""
        if index < len(self.prices):
            return self.prices[index], self.dates[index]
        return None, None


def build_histories(purchases, collapse_duplicates=True):
    """Per item, keep the `MAX_HISTORY` most recent purchases, newest first.

    `collapse_duplicates` drops repeated (date, price) pairs -- the same invoice
    line exported twice -- so the three slots hold three distinct purchases.
    """
    histories = {}
    ordered = purchases.sort_values(
        ["date", "source_row"], ascending=[False, False], na_position="last"
    )
    for key, group in ordered.groupby("key", sort=False):
        rows = group.dropna(subset=["date", "price"])
        if collapse_duplicates:
            rows = rows.drop_duplicates(subset=["date", "price"])
        head = rows.head(MAX_HISTORY)
        described = group["description"].loc[group["description"] != ""]
        histories[key] = ItemHistory(
            item=group["item"].iloc[0],
            description=described.iloc[0] if len(described) else "",
            prices=[round(float(price), 4) for price in head["price"]],
            dates=[when.to_pydatetime() for when in head["date"]],
            total_purchases=len(rows),
        )
    return histories


# --- writing the Order workbook -------------------------------------------


def _find_header_row(worksheet, wanted="itemcode", limit=10):
    for row in range(1, min(limit, worksheet.max_row) + 1):
        for col in range(1, worksheet.max_column + 1):
            if norm(worksheet.cell(row, col).value) == wanted:
                return row
    raise BuildError("No 'Item Code' header found in the Order workbook.")


def _map_order_columns(worksheet, header_row):
    labels = {}
    for col in range(1, worksheet.max_column + 1):
        value = worksheet.cell(header_row, col).value
        if value is not None:
            labels[value] = col
    resolved = _resolve(labels, ORDER_TARGETS)
    return {logical: labels[label] for logical, label in resolved.items()}


def _find_total_row(worksheet, header_row):
    for row in range(worksheet.max_row, header_row, -1):
        for col in range(1, worksheet.max_column + 1):
            if norm(worksheet.cell(row, col).value).startswith("total"):
                return row
    return None


def _capture_row(worksheet, row):
    """Snapshot a prototype row's styles and formulas so it can be replayed."""
    return {
        "height": worksheet.row_dimensions[row].height,
        "cells": {
            col: (copy(worksheet.cell(row, col)._style), worksheet.cell(row, col).value)
            for col in range(1, worksheet.max_column + 1)
        },
    }


def _rewrite_sum(formula, first_row, last_row):
    """Point a total-row SUM at the new data block: SUM(U2:U96) -> SUM(U2:U31)."""
    return re.sub(
        r"(SUM\()([A-Z]{1,3})\d+:([A-Z]{1,3})\d+(\))",
        lambda m: f"{m.group(1)}{m.group(2)}{first_row}:{m.group(3)}{last_row}{m.group(4)}",
        formula,
        flags=re.IGNORECASE,
    )


def _apply_prototype(worksheet, row, prototype, origin_row):
    """Lay a captured data row onto `row`, translating its formulas."""
    for col, (style, value) in prototype["cells"].items():
        cell = worksheet.cell(row, col)
        cell._style = copy(style)
        if isinstance(value, str) and value.startswith("="):
            cell.value = Translator(
                value, origin=f"{get_column_letter(col)}{origin_row}"
            ).translate_formula(f"{get_column_letter(col)}{row}")
        else:
            cell.value = None
    if prototype["height"]:
        worksheet.row_dimensions[row].height = prototype["height"]


def _history_formats(worksheet, row, columns):
    """Number formats for the six history cells, taken from slot 1.

    The supplied template formats `Last U.P.(3)` as a date, which would render
    the third price as one; slot 1 is the reference and the constants are the
    fallback when it is unusable too.
    """
    price_column = columns.get("price1")
    date_column = columns.get("date1")
    price_format = (
        worksheet.cell(row, price_column).number_format if price_column else PRICE_FORMAT
    )
    date_format = (
        worksheet.cell(row, date_column).number_format if date_column else DATE_FORMAT
    )
    if is_date_format(price_format):
        price_format = PRICE_FORMAT
    if not is_date_format(date_format):
        date_format = DATE_FORMAT
    return price_format, date_format


def _write_history(worksheet, row, columns, history):
    """Fill the six Last U.P. / Date cells for one item, newest price first."""
    price_format, date_format = _history_formats(worksheet, row, columns)
    for slot in range(MAX_HISTORY):
        price, when = history.cell(slot)
        price_col = columns.get(f"price{slot + 1}")
        date_col = columns.get(f"date{slot + 1}")
        if price_col:
            cell = worksheet.cell(row, price_col)
            cell.number_format = price_format
            if price is not None:
                cell.value = price
        if date_col:
            cell = worksheet.cell(row, date_col)
            cell.number_format = date_format
            if when is not None:
                cell.value = when


def build_order(histories, template=None, number_column=1):
    """Render the Order workbook from the template, one row per item."""
    workbook = load_workbook(template or TEMPLATE_PATH)
    worksheet = workbook.active

    header_row = _find_header_row(worksheet)
    columns = _map_order_columns(worksheet, header_row)
    if "item" not in columns:
        raise BuildError("The Order template has no 'Item Code' column.")

    total_row = _find_total_row(worksheet, header_row)
    first_data_row = header_row + 1
    proto_data = _capture_row(worksheet, first_data_row)
    proto_total = _capture_row(worksheet, total_row) if total_row else None

    last_template_row = total_row or worksheet.max_row
    if last_template_row >= first_data_row:
        worksheet.delete_rows(first_data_row, last_template_row - first_data_row + 1)

    items = sorted(histories.values(), key=lambda history: history.item)
    for offset, history in enumerate(items):
        row = first_data_row + offset
        _apply_prototype(worksheet, row, proto_data, first_data_row)
        if number_column:
            worksheet.cell(row, number_column).value = offset + 1
        worksheet.cell(row, columns["item"]).value = history.item
        if "description" in columns:
            worksheet.cell(row, columns["description"]).value = history.description
        _write_history(worksheet, row, columns, history)

    if proto_total:
        row = first_data_row + len(items)
        last_data_row = max(first_data_row + len(items) - 1, first_data_row)
        worksheet.row_dimensions[row].height = proto_total["height"]
        for col, (style, value) in proto_total["cells"].items():
            cell = worksheet.cell(row, col)
            cell._style = copy(style)
            cell.value = (
                _rewrite_sum(value, first_data_row, last_data_row)
                if isinstance(value, str) and value.startswith("=")
                else value
            )

    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def fill_existing_order(order_source, histories):
    """Fill the six history columns of an Order file that already lists items.

    Returns (workbook bytes, matched item codes, unmatched item codes).
    """
    workbook = load_workbook(order_source)
    worksheet = workbook.active

    header_row = _find_header_row(worksheet)
    columns = _map_order_columns(worksheet, header_row)
    if "item" not in columns:
        raise BuildError("The Order file has no 'Item Code' column.")

    matched, unmatched = [], []
    for row in range(header_row + 1, worksheet.max_row + 1):
        raw = worksheet.cell(row, columns["item"]).value
        code = str(raw).strip() if raw is not None else ""
        if not code or norm(code).startswith("total"):
            continue
        history = histories.get(code.upper())
        if history is None:
            unmatched.append(code)
            continue
        _write_history(worksheet, row, columns, history)
        matched.append(code)

    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue(), matched, unmatched


def histories_to_frame(histories):
    """Flatten histories into the preview table shown in the app."""
    rows = []
    for history in sorted(histories.values(), key=lambda item: item.item):
        row = {"Item Code": history.item, "Description": history.description}
        for slot in range(MAX_HISTORY):
            price, when = history.cell(slot)
            row[f"Last U.P.({slot + 1}) (EUR)"] = price
            row[f"Date({slot + 1})"] = when.date() if when else None
        row["Purchases on file"] = history.total_purchases
        rows.append(row)
    return pd.DataFrame(rows)


def main(argv=None):
    import argparse

    parser = argparse.ArgumentParser(description="Build the Order workbook.")
    parser.add_argument("purchases", help="last purchases workbook (.xlsx/.csv)")
    parser.add_argument("-o", "--output", default="Order (generated).xlsx")
    parser.add_argument("--template", default=None, help="Order template to use")
    parser.add_argument("--into", default=None, help="fill this existing Order file")
    parser.add_argument(
        "--keep-duplicates",
        action="store_true",
        help="keep repeated (date, price) rows instead of collapsing them",
    )
    args = parser.parse_args(argv)

    histories = build_histories(
        read_purchases(args.purchases), collapse_duplicates=not args.keep_duplicates
    )
    if args.into:
        data, matched, unmatched = fill_existing_order(args.into, histories)
        print(f"Filled {len(matched)} item(s); {len(unmatched)} without history.")
    else:
        data = build_order(histories, template=args.template)
        print(f"Wrote {len(histories)} item(s).")
    Path(args.output).write_bytes(data)
    print(f"Saved {args.output}")


if __name__ == "__main__":
    main()
